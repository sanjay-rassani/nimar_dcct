"""Excel export for Delivery Points."""
from __future__ import annotations

from io import BytesIO

from django.db.models import Q
from openpyxl import Workbook
from openpyxl.styles import Font

from tracker.models import Issue, IssueStatus, Point, PointStatus

EXPORT_HEADERS = [
    "Point ID",
    "Annexure",
    "Number",
    "Module",
    "Phase",
    "Title",
    "Requirement",
    "Status",
    "Forbmax Response",
    "NIMAR Response",
    "Decision",
    "Remarks",
    "Owner",
    "Due Date",
    "Open Issues",
    "Updated At",
    "Updated By",
]


def filtered_points_queryset(params):
    qs = Point.objects.select_related("module", "phase", "owner", "updated_by").order_by("external_id")

    q = (params.get("q") or "").strip()
    if q:
        qs = qs.filter(
            Q(external_id__icontains=q)
            | Q(title__icontains=q)
            | Q(requirement__icontains=q)
            | Q(module__name__icontains=q)
            | Q(remarks__icontains=q)
        ).distinct()

    if phase := params.get("phase"):
        if str(phase).isdigit():
            qs = qs.filter(phase_id=int(phase))
        else:
            qs = qs.filter(phase__code=phase)
    if module := params.get("module"):
        if str(module).isdigit():
            qs = qs.filter(module_id=int(module))
        else:
            qs = qs.filter(module__name=module)
    if status := params.get("status"):
        qs = qs.filter(status=status)
    if annexure := params.get("annexure"):
        qs = qs.filter(annexure=annexure)
    if owner := params.get("owner"):
        if str(owner).isdigit():
            qs = qs.filter(owner_id=int(owner))
    if issue_status := params.get("issue_status"):
        qs = qs.filter(issues__status=issue_status).distinct()
    return qs


def export_points_xlsx(params) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Points"
    ws.append(EXPORT_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    qs = filtered_points_queryset(params)
    for point in qs:
        open_issues = point.issues.exclude(status=IssueStatus.CLOSED).count()
        ws.append(
            [
                point.external_id,
                point.annexure,
                point.number,
                point.module.name if point.module_id else "",
                point.phase.name if point.phase_id else "",
                point.title,
                point.requirement,
                PointStatus(point.status).label if point.status in PointStatus.values else point.status,
                point.forbmax_response,
                point.nimar_response,
                point.decision,
                point.remarks,
                point.owner.name if point.owner_id else "",
                point.due_date.isoformat() if point.due_date else "",
                open_issues,
                point.updated_at.isoformat() if point.updated_at else "",
                point.updated_by.name if point.updated_by_id else "",
            ]
        )

    # Optional issues sheet when requested
    if params.get("include_issues") in {"1", "true", "yes"}:
        ws2 = wb.create_sheet("Issues")
        ws2.append(
            [
                "Issue ID",
                "Point ID",
                "Title",
                "Description",
                "Status",
                "Raised By",
                "Assigned To",
                "Resolved At",
                "Created At",
            ]
        )
        for cell in ws2[1]:
            cell.font = Font(bold=True)
        issues = Issue.objects.select_related("point", "raised_by", "assigned_to").filter(
            point__in=qs
        )
        if issue_status := params.get("issue_status"):
            issues = issues.filter(status=issue_status)
        for issue in issues.order_by("point__external_id", "id"):
            ws2.append(
                [
                    issue.pk,
                    issue.point.external_id,
                    issue.title,
                    issue.description,
                    IssueStatus(issue.status).label,
                    issue.raised_by.name if issue.raised_by_id else "",
                    issue.assigned_to.name if issue.assigned_to_id else "",
                    issue.resolved_at.isoformat() if issue.resolved_at else "",
                    issue.created_at.isoformat() if issue.created_at else "",
                ]
            )

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
